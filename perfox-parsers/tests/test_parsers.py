"""Tests for the Perfox resume parser service (perfox-parsers/).

    python -m pytest perfox-parsers/tests -p no:cacheprovider            offline: nothing leaves the machine
    RUN_LIVE=1 python -m pytest perfox-parsers/tests -k live             also the real Perfox call (fake resume only)

The five sample resumes are built here, in pytest's temp folder, on every run, so no resume file lives in the repository.
Every person, e-mail address and phone number in them is invented. Offline tests never reach Perfox: call_perfox is
replaced, and any test that forgets to do so fails loudly instead of calling out.
"""
import io
import logging
import os
import sys
from datetime import date
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))  # parsers.py, schema.py, fix_normlization.py sit next to tests/

import parsers  # noqa: E402  (parsers.py puts the repo root on sys.path itself)
import masking  # noqa: E402
from fix_normlization import check_required, fix_dates_and_experience, phone_numbers, tidy_education  # noqa: E402
from schema import build_schema  # noqa: E402
from verify_identity import verify_identity  # noqa: E402

PDF, DOCX = "application/pdf", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# Invented identities. The privacy test searches the logs for every one of these strings.
NAME, EMAIL, PHONE = "Aarav Sharma", "aarav.sharma@example.com", "+91 98765 43210"
NO_EMAIL_NAME, NO_EMAIL_PHONE = "Priya Verma", "+91 91234 56789"
EMPLOYERS = ["Zenithwave Technologies", "Brightloom Systems", "Quillstone Labs"]


# ---------- the five sample resumes ----------

def _pdf(lines: list[str]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setTitle("Resume")
    c.setAuthor("")
    y = 800
    for line in lines:
        c.setFont("Helvetica-Bold" if line.isupper() else "Helvetica", 11)
        c.drawString(56, y, line)
        y -= 20
    c.save()
    return buf.getvalue()


def _docx(lines: list[str]) -> bytes:
    import docx
    d = docx.Document()
    for line in lines:
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _scanned_pdf() -> bytes:
    """A picture of a resume on a page, with no text layer."""
    from PIL import Image, ImageDraw
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    img = Image.new("RGB", (1240, 1754), "white")
    d = ImageDraw.Draw(img)
    for i, t in enumerate(["Meera Nair", "EXPERIENCE", "Analyst at Quillstone Labs 2019 - 2024", "EDUCATION", "B.Com 2019"]):
        d.text((100, 120 + i * 60), t, fill="black")
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.drawImage(ImageReader(img), 0, 0, width=595, height=842)
    c.save()
    return buf.getvalue()


NORMAL_LINES = [
    NAME, f"{EMAIL} | {PHONE} | Bengaluru, India", "SUMMARY",
    "Backend engineer with six years of experience building payment and logistics services in Python.",
    "EXPERIENCE", f"Senior Engineer, {EMPLOYERS[0]}  (Jan 2021 - Present)",
    "Built the order tracking service used by forty warehouses.",
    f"Software Engineer, {EMPLOYERS[1]}  (Jul 2018 - Dec 2020)", "Moved the billing batch jobs to a queue based design.",
    "EDUCATION", "B.Tech in Computer Science, Riverbend Institute of Technology, 2018", "SKILLS", "Python, FastAPI, PostgreSQL, Docker",
]
NO_EMAIL_LINES = [
    NO_EMAIL_NAME, f"{NO_EMAIL_PHONE} | Pune, India", "EXPERIENCE", f"Data Analyst, {EMPLOYERS[2]}  (Mar 2019 - Present)",
    "Prepared weekly sales dashboards and cleaned the regional data sets for the finance team.",
    "EDUCATION", "B.Com, Lakeview College, 2018", "SKILLS", "Excel, SQL, Power BI",
]
OVERLAP_LINES = [
    "Karan Malhotra", "+91 99887 76655 | karan.malhotra@example.com", "EXPERIENCE",
    f"Contract Developer, {EMPLOYERS[0]}  (Jan 2018 - Dec 2019)",
    f"Freelance Consultant, {EMPLOYERS[1]}  (Jun 2019 - Jun 2020)",
    f"Lead Developer, {EMPLOYERS[2]}  (Jan 2022 - Present)",
    "EDUCATION", "B.E. in Electronics, Harborview University, 2017", "SKILLS", "Java, Spring, Kafka",
]


@pytest.fixture(scope="session")
def samples(tmp_path_factory) -> dict[str, Path]:
    folder = tmp_path_factory.mktemp("samples")
    files = {"normal_resume.pdf": _pdf(NORMAL_LINES), "normal_resume.docx": _docx(NORMAL_LINES),
             "no_email_resume.pdf": _pdf(NO_EMAIL_LINES), "overlapping_jobs_resume.pdf": _pdf(OVERLAP_LINES),
             "scanned_resume.pdf": _scanned_pdf()}
    for name, data in files.items():
        (folder / name).write_bytes(data)
    return {name: folder / name for name in files}


# ---------- plumbing ----------

@pytest.fixture
def client():
    return TestClient(parsers.app)


@pytest.fixture(autouse=True)
def no_real_perfox(request, monkeypatch):
    """Offline tests must never reach Perfox, whatever the .env says. The live test is exempt."""
    if request.node.name.startswith("test_live"):
        return

    async def refuse(*args, **kwargs):
        raise AssertionError("this offline test tried to call Perfox")
    monkeypatch.setattr(parsers, "call_perfox", refuse)


USAGE = {"input_tokens": 500, "output_tokens": 400}


def fake_perfox(monkeypatch, parsed: dict, usage=USAGE):
    async def answer(*args):
        return parsed, usage, None
    monkeypatch.setattr(parsers, "call_perfox", answer)


def post(client, name: str, data: bytes, mime: str = PDF):
    return client.post("/parsing/resume", files={"pdf_doc": (name, data, mime)})


def parsed_resume(**over) -> dict:
    """What Perfox would return: every schema key present, as the schema demands."""
    out = {k: ([] if v.get("type") == "array" else {} if v.get("type") == "object" else "")
           for k, v in build_schema()["properties"].items()}
    out.update({"name": NAME, "email": EMAIL, "phoneNumber": PHONE, "skills": ["Python"]})
    out.update(over)
    return out


# ---------- intake: 400 ----------

def test_empty_file_is_400(client):
    r = post(client, "empty.pdf", b"")
    assert r.status_code == 400
    assert r.json()["error"]["code"] == 400
    assert r.json()["error"]["error_type"] == "INVALID_FILE"
    assert r.json()["success"] is False and r.json()["error"]["request_id"]
    assert r.headers.get("X-Request-ID") == r.json()["error"]["request_id"]


def test_oversized_file_is_400(client):
    r = post(client, "big.pdf", b"%PDF-1.4\n" + b"0" * parsers.MAX_FILE_BYTES)
    assert r.status_code == 400
    assert r.json()["error"]["code"] == 400
    assert r.json()["error"]["error_type"] == "INVALID_FILE"
    assert "too large" in r.json()["error"]["message"]
    assert r.headers.get("X-Request-ID")


def test_exactly_at_the_limit_is_not_rejected_for_size(client):
    r = post(client, "limit.pdf", b"0" * parsers.MAX_FILE_BYTES)  # not a real PDF, so it fails later, on reading, as 422
    assert r.status_code == 422
    assert r.json()["error"]["code"] == 422
    assert r.json()["error"]["error_type"] == "UNREADABLE_DOCUMENT"


def test_unsupported_extension_is_400(client):
    r = post(client, "resume.txt", b"plain text resume " * 10, "text/plain")
    assert r.status_code == 400
    assert r.json()["error"]["code"] == 400
    assert r.json()["error"]["error_type"] == "INVALID_FILE"
    assert ".txt" in r.json()["error"]["message"]
    assert r.headers.get("X-Request-ID")


# ---------- extraction: 422 ----------

def test_scanned_pdf_is_unreadable(client, samples):
    r = post(client, "scanned_resume.pdf", samples["scanned_resume.pdf"].read_bytes())
    assert r.status_code == 422
    assert r.json()["error"]["code"] == 422
    assert r.json()["error"]["error_type"] == "UNREADABLE_DOCUMENT"
    assert "scanned" in r.json()["error"]["message"]
    assert r.headers.get("X-Request-ID")


def test_corrupt_pdf_is_unreadable(client):
    r = post(client, "broken.pdf", b"this is not a pdf at all " * 20)
    assert r.status_code == 422
    assert r.json()["error"]["code"] == 422
    assert r.json()["error"]["error_type"] == "UNREADABLE_DOCUMENT"
    assert r.headers.get("X-Request-ID")


def test_resume_without_email_is_missing_required_fields(client, samples, monkeypatch):
    raw = samples["no_email_resume.pdf"].read_bytes()
    text, _, _ = parsers.get_text(raw, ".pdf")
    from pypdf import PdfReader
    meta = " ".join(str(v) for v in (PdfReader(io.BytesIO(raw)).metadata or {}).values())
    assert "@" not in text and "@" not in meta  # the sample really has no e-mail: not in the text, not in the document properties
    fake_perfox(monkeypatch, parsed_resume(name=NO_EMAIL_NAME, email="", phoneNumber=NO_EMAIL_PHONE))
    r = post(client, "no_email_resume.pdf", raw)
    assert r.status_code == 422
    assert r.json()["error"]["code"] == 422
    assert r.json()["error"]["error_type"] == "MISSING_REQUIRED_FIELDS"
    assert r.json()["error"]["message"] == (
        "We couldn't find the following details in this resume: email. "
        "Please make sure it clearly shows the name, email, and phone number."
    )
    assert r.headers.get("X-Request-ID")


# ---------- the whole route, Perfox faked ----------

@pytest.mark.parametrize("name,mime", [("normal_resume.pdf", PDF), ("normal_resume.docx", DOCX)])
def test_normal_resume_returns_the_parsed_data(client, samples, monkeypatch, name, mime):
    fake_perfox(monkeypatch, parsed_resume())
    r = post(client, name, samples[name].read_bytes(), mime)
    assert r.status_code == 200
    body = r.json()
    assert "success" not in body and "data" not in body
    assert body["email"] == EMAIL and "total_experience" in body
    assert r.headers.get("X-Request-ID")
    assert list(body)[-1] == "usage"


def test_overlapping_and_ongoing_jobs_through_the_route(client, samples, monkeypatch):
    jobs = [{"jobTitle": "Contract Developer", "startDate": "2018-01-01", "endDate": "2019-12-31", "isCurrent": False},
            {"jobTitle": "Freelance Consultant", "startDate": "2019-06-01", "endDate": "2020-06-30", "isCurrent": False},
            {"jobTitle": "Lead Developer", "startDate": "2022-01-01", "endDate": "2022-01-01", "isCurrent": True}]
    fake_perfox(monkeypatch, parsed_resume(workExperience=jobs))
    r = post(client, "overlapping_jobs_resume.pdf", samples["overlapping_jobs_resume.pdf"].read_bytes())
    assert r.status_code == 200
    out = r.json()
    assert "data" not in out
    today = date.today()
    assert out["workExperience"][2]["endDate"] == today.isoformat()  # "Present" is today, not the start month
    from fix_normlization import _years_months
    expected = _years_months((date(2020, 6, 30) - date(2018, 1, 1)).days + (today - date(2022, 1, 1)).days)
    assert out["total_experience"] == expected


# ---------- token usage ----------

def test_the_response_carries_the_token_counts(client, samples, monkeypatch):
    fake_perfox(monkeypatch, parsed_resume(), {"input_tokens": 527, "output_tokens": 481})
    r = post(client, "normal_resume.pdf", samples["normal_resume.pdf"].read_bytes())
    assert r.json()["usage"] == {"input_tokens": 527, "output_tokens": 481, "total_tokens": 1008}
    assert list(r.json())[-1] == "usage"  # last, so it sits at the bottom of the output


@pytest.mark.parametrize("usage", [None, {}, {"input_tokens": 5}, {"input_tokens": "x", "output_tokens": 1}])
def test_missing_or_broken_usage_does_not_break_the_parse(client, samples, monkeypatch, usage):
    fake_perfox(monkeypatch, parsed_resume(), usage)
    r = post(client, "normal_resume.pdf", samples["normal_resume.pdf"].read_bytes())
    assert r.status_code == 200 and r.json()["usage"] is None


# ---------- fix_dates_and_experience ----------

TODAY = date(2026, 6, 1)


def months(d: dict) -> int:
    return d["years"] * 12 + d["months"]


def test_overlapping_roles_are_not_counted_twice():
    parsed = {"workExperience": [
        {"startDate": "2020-01-01", "endDate": "2022-12-31"},
        {"startDate": "2021-06-01", "endDate": "2023-06-30"},  # runs inside and past the first
        {"startDate": "2024-01-01", "isCurrent": True}]}
    out = fix_dates_and_experience(parsed, TODAY)
    first, second, current = (j["duration"] for j in out["workExperience"])
    assert out["total_experience"] == {"years": 5, "months": 11}  # 2020-01 to 2023-06, then 2024-01 to 2026-06
    assert months(first) + months(second) + months(current) == 90  # what naive adding would claim: 7y 6m
    assert months(out["total_experience"]) < months(first) + months(second) + months(current)


def test_a_role_inside_another_adds_nothing():
    parsed = {"workExperience": [{"startDate": "2019-01-01", "endDate": "2023-01-01"},
                                 {"startDate": "2020-01-01", "endDate": "2021-01-01"}]}
    total = fix_dates_and_experience(parsed, TODAY)["total_experience"]
    assert total == {"years": 4, "months": 0}


def test_ongoing_role_ends_today_and_future_end_dates_are_clamped():
    parsed = {"workExperience": [{"startDate": "2025-06-01", "endDate": "2025-06-01", "isCurrent": True},
                                 {"startDate": "2024-01-01", "endDate": "2031-01-01"}]}
    jobs = fix_dates_and_experience(parsed, TODAY)["workExperience"]
    assert jobs[0]["endDate"] == "2026-06-01" and jobs[0]["duration"] == {"years": 1, "months": 0}
    assert months(jobs[1]["duration"]) == 29  # 2024-01-01 to today, not to 2031


def test_no_jobs_means_zero_experience_and_does_not_crash():
    assert fix_dates_and_experience({"workExperience": []}, TODAY)["total_experience"] == {"years": 0, "months": 0}
    assert fix_dates_and_experience({"workExperience": [{"startDate": "not a date"}]}, TODAY)["total_experience"] == {"years": 0, "months": 0}


def test_certification_names_are_removed_from_technical_skills():
    parsed = {"certifications": [{"certificationName": "AWS Certified Developer"}],
              "technical_skills": ["Python", "aws certified developer", "Docker"]}
    assert fix_dates_and_experience(parsed, TODAY)["technical_skills"] == ["Python", "Docker"]


# ---------- the regex rescue ----------

def test_regex_fills_in_email_and_phone_the_model_left_out():
    text = "Rohan Gupta\nrohan.gupta@example.com | 98765 43210 | Delhi\nEXPERIENCE ..."
    parsed, missing = check_required({"name": "Rohan Gupta", "email": "", "phoneNumber": ""}, text)
    assert missing == []
    assert parsed["email"] == "rohan.gupta@example.com" and parsed["phoneNumber"] == "98765 43210"


def test_regex_replaces_a_malformed_email_but_keeps_a_good_one():
    text = "contact: real.one@example.com"
    assert check_required({"name": "A", "email": "not-an-email", "phoneNumber": "9876543210"}, text)[0]["email"] == "real.one@example.com"
    assert check_required({"name": "A", "email": "kept@example.com", "phoneNumber": "9876543210"}, text)[0]["email"] == "kept@example.com"


def test_regex_cannot_invent_what_the_text_does_not_have():
    parsed, missing = check_required({"name": "", "email": "", "phoneNumber": ""}, "just some words with no contact details")
    assert missing == ["name", "email", "phoneNumber"]


@pytest.mark.parametrize("text", [
    "Analyst 2019 - 2023 - 2024 Pune",
    "Jan 2019 - Dec 2023\nMar 2015 - 2018",
    "Employed 2019 2020 2021 2022",
    "Roll no 12345 and 678",
    "Account 1234567890123456",
])
def test_dates_and_other_numbers_are_not_taken_for_a_phone_number(text):
    parsed, missing = check_required({"name": "A", "email": "a@example.com", "phoneNumber": ""}, text)
    assert missing == ["phoneNumber"] and not parsed["phoneNumber"]


@pytest.mark.parametrize("text,expected", [
    ("+91 98765 43210", "+91 98765 43210"),
    ("Call 98765 43210 anytime", "98765 43210"),
    ("9876543210", "9876543210"),
    ("+91-9876543210 | Pune", "+91-9876543210"),
    ("(022) 2345 6789", "(022) 2345 6789"),
    ("Analyst 2019 - 2023 | +91 98765 43210", "+91 98765 43210"),  # the date range comes first, the real number is still found
    ("2019 - 2023\n98765 43210", "98765 43210"),                    # digits on the next line are not joined on
])
def test_real_phone_formats_are_still_found(text, expected):
    assert check_required({"name": "A", "email": "a@example.com", "phoneNumber": ""}, text)[0]["phoneNumber"] == expected


# ---------- privacy: nothing from the resume in the logs ----------

class _Capture(logging.Handler):
    def __init__(self):
        super().__init__(logging.DEBUG)
        self.lines: list[str] = []

    def emit(self, record):
        self.lines.append(record.getMessage())


def test_no_resume_text_or_personal_data_reaches_the_logs(client, samples, monkeypatch):
    cap = _Capture()
    watched = [logging.getLogger(), parsers.log, logging.getLogger("httpx")]
    old = [(lg, lg.level) for lg in watched]
    for lg in watched:
        lg.addHandler(cap)
        lg.setLevel(logging.DEBUG)
    try:
        fake_perfox(monkeypatch, parsed_resume())
        post(client, "normal_resume.pdf", samples["normal_resume.pdf"].read_bytes())          # success
        post(client, "normal_resume.docx", samples["normal_resume.docx"].read_bytes(), DOCX)  # success
        post(client, "scanned_resume.pdf", samples["scanned_resume.pdf"].read_bytes())        # 422 unreadable
        post(client, "resume.txt", b"x" * 100, "text/plain")                                  # 400
        fake_perfox(monkeypatch, parsed_resume(name=NO_EMAIL_NAME, email="", phoneNumber=NO_EMAIL_PHONE))
        post(client, "no_email_resume.pdf", samples["no_email_resume.pdf"].read_bytes())      # 422 missing fields

        async def refused(*args):
            return None, None, (502, "PERFOX_REJECTED", "Perfox rejected the request: Bad Request")
        monkeypatch.setattr(parsers, "call_perfox", refused)
        post(client, "normal_resume.pdf", samples["normal_resume.pdf"].read_bytes())          # upstream failure
    finally:
        for lg, level in old:
            lg.removeHandler(cap)
            lg.setLevel(level)

    logged = "\n".join(cap.lines)
    assert len(cap.lines) >= 6  # something was logged, so the check below means something
    secrets = [NAME, EMAIL, PHONE, NO_EMAIL_NAME, NO_EMAIL_PHONE, "98765", "Bengaluru", "payment and logistics", "Riverbend",
               "Power BI", "Lakeview", *EMPLOYERS]
    leaked = [s for s in secrets if s.lower() in logged.lower()]
    assert not leaked, f"found in the logs: {leaked}"


# ---------- live: the real Perfox endpoint, a fake resume only ----------

@pytest.mark.skipif(os.getenv("RUN_LIVE") != "1", reason="set RUN_LIVE=1 to call the real Perfox endpoint")
@pytest.mark.skipif(not (parsers.PERFOX_BASE_URL and parsers.PERFOX_API_KEY), reason="PERFOX_BASE_URL / PERFOX_API_KEY not set")
def test_live_perfox_parses_the_fake_resume(client, samples):
    r = post(client, "normal_resume.pdf", samples["normal_resume.pdf"].read_bytes())
    assert r.status_code == 200, r.text[:300]
    body = r.json()
    assert "success" not in body and "data" not in body  # flat, like Questlight's
    data, schema = body, build_schema()
    assert set(schema["required"]) <= set(data), f"keys missing from the answer: {set(schema['required']) - set(data)}"
    assert data["name"].strip() and "@" in data["email"] and data["phoneNumber"].strip()
    assert set(data["total_experience"]) == {"years", "months"}
    usage = body["usage"]
    assert usage["input_tokens"] > 0 and usage["output_tokens"] > 0
    assert usage["total_tokens"] == usage["input_tokens"] + usage["output_tokens"]


def test_jd_schema_has_all_25_required_fields():
    from jd_schema import build_jd_schema

    schema = build_jd_schema()
    assert schema["type"] == "object"
    assert schema["additionalProperties"] is False
    assert len(schema["properties"]) == 25
    # Perfox silently drops any field not in required:
    assert schema["required"] == list(schema["properties"])
    assert "salaryAsWritten" in schema["properties"]
    assert "experienceAsWritten" in schema["properties"]


def test_build_jd_prompt_interpolates_today():
    from datetime import date
    from jd_schema import build_jd_prompt

    today = date(2026, 10, 19)
    prompt = build_jd_prompt(today)
    assert "Today's date is 2026-10-19." in prompt
    assert "salaryAsWritten" in prompt
    assert "experienceAsWritten" in prompt


def test_fix_jd_experience_salary_and_normalization():
    from datetime import date
    from fix_jd import fix_jd, parse_experience, parse_salary

    # 1. Salary tests (BUILD_SPEC requirements)
    assert parse_salary("12-18 LPA") == ("1200000", "1800000")
    assert parse_salary("80k-1L") == ("80000", "100000")
    assert parse_salary("18-12 LPA") == ("1200000", "1800000")  # reversed range swapped
    assert parse_salary("50,000 per month") == ("600000", "600000")  # monthly * 12
    assert parse_salary("20 LPA") == ("2000000", "2000000")  # single figure used for both
    assert parse_salary("") == ("", "")  # missing value
    assert parse_salary(None) == ("", "")

    # 2. Experience tests (BUILD_SPEC requirements)
    assert parse_experience("4-8 years") == (4, 8)
    assert parse_experience("8 to 4 years") == (4, 8)  # reversed range swapped
    assert parse_experience("5+ years") == (5, "")
    assert parse_experience("") == ("", "")  # missing value

    # 3. Full 30-field contract & normalization test
    today = date(2026, 10, 19)
    sample_model_output = {
        "jobPositionTitle": "  Senior QA Automation Engineer  ",
        "jobSummary": "Lead automation testing.",
        "comments": "Strong Cypress preferred.",
        "customerName": "Acme Corp",
        "country": "India", "state": "Karnataka", "city": "Bengaluru", "workAddress": "",
        "industry": "IT", "domain": "Software Testing",
        "primarySkills": ["Java", "Selenium WebDriver"],
        "skillsRequired": ["Selenium WebDriver", "TestNG", "JIRA"],
        "skillDomains": ["QA Automation"],
        "degree": ["B.Tech"],
        "languages": ["English"],
        "jobType": "Full time",
        "workMode": "remote friendly",
        "noticePeriodMonths": "30 days",
        "numberOfOpenings": "2",
        "closingDate": "2026-12-31",
        "jobDurationMonths": "",
        "jobWorkHoursPerDay": "8",
        "priority": "High priority",
        "salaryAsWritten": "12-18 LPA",
        "experienceAsWritten": "4-8 years",
    }
    fixed = fix_jd(sample_model_output, today)

    assert len(fixed) == 30
    assert "salaryAsWritten" not in fixed
    assert "experienceAsWritten" not in fixed
    assert fixed["minExperienceYears"] == 4
    assert fixed["maxExperienceYears"] == 8
    assert fixed["minimumSalary"] == "1200000"
    assert fixed["maximumSalary"] == "1800000"
    assert fixed["jobType"] == "Full-time"
    assert fixed["workMode"] == "Remote"
    assert fixed["priority"] == "High"
    assert fixed["noticePeriodMonths"] == "1"
    assert fixed["numberOfOpenings"] == 2
    assert fixed["postingDate"] == "2026-10-19"
    assert fixed["jobStatus"] == "Active"
    assert fixed["isOpenForReferral"] is False
    # Primary skills must be automatically included in skillsRequired
    assert "Java" in fixed["skillsRequired"]


# ── JD Route Tests ───────────────────────────────────────────────────────────

def parsed_jd(**over) -> dict:
    """Mock Perfox response for JD extraction containing all 25 schema keys."""
    from jd_schema import build_jd_schema
    schema = build_jd_schema()
    out = {
        k: ([] if v.get("type") == "array" else "")
        for k, v in schema["properties"].items()
    }
    out.update({
        "jobPositionTitle": "Senior QA Automation Engineer",
        "jobSummary": "Lead automation testing of web and API platforms.",
        "comments": "Strong Cypress and JMeter experience preferred.",
        "customerName": "Northwind Labs",
        "country": "India", "state": "Karnataka", "city": "Bengaluru",
        "primarySkills": ["Java", "Selenium"],
        "skillsRequired": ["Selenium", "TestNG", "JIRA"],
        "salaryAsWritten": "12-18 LPA",
        "experienceAsWritten": "4-8 years",
        "workMode": "Hybrid",
        "jobType": "Full-time",
    })
    out.update(over)
    return out


def post_jd(client, file=None, jd_text: str = ""):
    data = {}
    files = {}
    if jd_text:
        data["jd_text"] = jd_text
    if file:
        files["file"] = file
    return client.post("/parsing/jd", data=data, files=files)


def test_jd_no_input_is_400(client):
    r = post_jd(client)
    assert r.status_code == 400
    assert r.json()["error"]["code"] == 400
    assert r.json()["error"]["error_type"] == "NO_JOB_DESCRIPTION"
    assert r.headers.get("X-Request-ID")


def test_jd_unsupported_file_extension_is_400(client):
    r = post_jd(client, file=("jd.png", b"fake image bytes", "image/png"))
    assert r.status_code == 400
    assert r.json()["error"]["code"] == 400
    assert r.json()["error"]["error_type"] == "INVALID_FILE"
    assert r.headers.get("X-Request-ID")


def test_jd_too_short_text_is_422(client):
    r = post_jd(client, jd_text="Too short JD")
    assert r.status_code == 422
    assert r.json()["error"]["code"] == 422
    assert r.json()["error"]["error_type"] == "UNREADABLE_DOCUMENT"
    assert r.headers.get("X-Request-ID")


def test_jd_pasted_text_success(client, monkeypatch):
    fake_perfox(monkeypatch, parsed_jd())
    text = "Need a Senior QA Automation Engineer with Java and Selenium experience. 4-8 years exp. 12-18 LPA."
    r = post_jd(client, jd_text=text)
    assert r.status_code == 200
    body = r.json()
    assert "success" not in body and "data" not in body
    assert body["jobPositionTitle"] == "Senior QA Automation Engineer"
    assert body["minimumSalary"] == "1200000"
    assert body["maximumSalary"] == "1800000"
    assert body["minExperienceYears"] == 4
    assert body["maxExperienceYears"] == 8
    assert list(body)[-1] == "usage"
    assert r.headers.get("X-Request-ID")


def test_jd_file_upload_pdf(client, monkeypatch):
    fake_perfox(monkeypatch, parsed_jd())
    pdf_bytes = _pdf([
        "Senior QA Automation Engineer",
        "Northwind Labs - Bengaluru, India",
        "We need an experienced engineer with Java, Selenium, and TestNG.",
        "Experience: 4-8 years. Salary: 12-18 LPA.",
    ])
    r = post_jd(client, file=("job_spec.pdf", pdf_bytes, PDF))
    assert r.status_code == 200
    body = r.json()
    assert "success" not in body and "data" not in body
    assert body["jobPositionTitle"] == "Senior QA Automation Engineer"
    assert r.headers.get("X-Request-ID")


def test_jd_file_upload_docx(client, monkeypatch):
    fake_perfox(monkeypatch, parsed_jd())
    docx_bytes = _docx([
        "Senior QA Automation Engineer",
        "Northwind Labs - Bengaluru, India",
        "Experience: 4-8 years. Salary: 12-18 LPA.",
    ])
    r = post_jd(client, file=("job_spec.docx", docx_bytes, DOCX))
    assert r.status_code == 200
    body = r.json()
    assert "success" not in body and "data" not in body
    assert body["jobPositionTitle"] == "Senior QA Automation Engineer"
    assert r.headers.get("X-Request-ID")


def test_jd_file_upload_txt(client, monkeypatch):
    fake_perfox(monkeypatch, parsed_jd())
    txt_bytes = b"Senior QA Automation Engineer needed. Java and Selenium. 4 to 8 years experience. Salary 12-18 LPA."
    r = post_jd(client, file=("job_spec.txt", txt_bytes, "text/plain"))
    assert r.status_code == 200
    body = r.json()
    assert "success" not in body and "data" not in body
    assert body["jobPositionTitle"] == "Senior QA Automation Engineer"
    assert r.headers.get("X-Request-ID")


def test_jd_file_wins_over_pasted_text(client, monkeypatch):
    seen = {}
    async def spy(name, data, ext, text, schema, prompt):
        seen["name"] = name
        return parsed_jd(), USAGE, None
    monkeypatch.setattr(parsers, "call_perfox", spy)

    txt_bytes = b"File JD content with enough characters to pass validation threshold easily."
    r = post_jd(client, file=("from_file.txt", txt_bytes, "text/plain"), jd_text="Pasted text content with enough characters.")
    assert r.status_code == 200
    assert seen["name"] == "from_file.txt"
    assert r.headers.get("X-Request-ID")



def test_jd_privacy_zero_leaks_in_logs(client, monkeypatch, caplog):
    invented_client = "SuperSecretHiringClient12345"
    invented_title = "UnicornSpecialistEngineer98765"
    fake_perfox(monkeypatch, parsed_jd(customerName=invented_client, jobPositionTitle=invented_title))

    with caplog.at_level(logging.INFO, logger="perfox"):
        r = post_jd(client, jd_text=f"Role for {invented_title} at {invented_client}. Experience 5+ years.")
        assert r.status_code == 200

    all_logs = " ".join(record.getMessage() for record in caplog.records)
    assert invented_client not in all_logs
    assert invented_title not in all_logs


# ── Response Helper Unit Tests (Questlight format) ───────────────────────────

def test_fail_helper_layout_and_header():
    import json
    from parsers import fail

    resp = fail(422, "MISSING_REQUIRED_FIELDS", "Missing email", "req-test-123")
    assert resp.status_code == 422
    assert resp.headers.get("X-Request-ID") == "req-test-123"

    body = json.loads(resp.body.decode("utf-8"))
    assert body["success"] is False
    assert body["error"] == {
        "code": 422,
        "error_type": "MISSING_REQUIRED_FIELDS",
        "message": "Missing email",
        "request_id": "req-test-123",
    }


def test_ok_helper_layout_and_header():
    import json
    from parsers import ok

    content = {"name": "Aarav Sharma", "email": "aarav@example.com"}
    resp = ok(content, "req-test-456")
    assert resp.status_code == 200
    assert resp.headers.get("X-Request-ID") == "req-test-456"

    body = json.loads(resp.body.decode("utf-8"))
    assert body == content


# ── App Contract Tests (Running output through main app's code) ──────────────

def test_contract_resume_missing_details_integration_with_app(client, samples, monkeypatch):
    """Verifies app/intake/parser_questlight and questlight.unreadable_items work directly with our error."""
    from app.intake import parser_questlight, questlight

    fake_perfox(monkeypatch, parsed_resume(name=NO_EMAIL_NAME, email="", phoneNumber=NO_EMAIL_PHONE))
    r = post(client, "no_email_resume.pdf", samples["no_email_resume.pdf"].read_bytes())
    assert r.status_code == 422

    err_msg = r.json()["error"]["message"]
    # App check 1: matches MISSING_MESSAGE
    assert err_msg.startswith(parser_questlight.MISSING_MESSAGE)

    # App check 2: unreadable_items parses the missing field
    items = questlight.unreadable_items(err_msg)
    assert any(item["id"] == "email" for item in items)


def test_contract_jd_output_passes_through_app_from_parsed(client, monkeypatch):
    """Verifies app/intake/jobs.from_parsed consumes our flat /parsing/jd output directly."""
    from app.intake.jobs import from_parsed

    fake_perfox(monkeypatch, parsed_jd())
    r = post_jd(client, jd_text="Need a Senior QA Automation Engineer with Java and Selenium. 4-8 years exp. 12-18 LPA.")
    assert r.status_code == 200

    body = r.json()
    assert list(body)[-1] == "usage"
    usage = body.pop("usage")
    assert usage["total_tokens"] > 0

    # Pass the flat parsed output directly to the app's from_parsed:
    job = from_parsed(body)
    assert job["title"] == "Senior QA Automation Engineer"
    assert job["salary_min"] == 1200000.0
    assert job["salary_max"] == 1800000.0
    assert job["min_exp"] == 4.0
    assert job["max_exp"] == 8.0
    assert "Java" in job["skills"]


# ---------- identity checks against the document's own text (verify_identity) ----------

ID_TEXT = ("Lakshmi Iyer\nlakshmi.iyer@example.com | +91 98450 12345 | Chennai\nEXPERIENCE\n"
           "Analyst at Quillstone Labs (Jan 2019 - Present)\nReferee: Jane Doe")


def ident(**over) -> dict:
    out = {"name": "Lakshmi Iyer", "email": "lakshmi.iyer@example.com", "phoneNumber": "+91 98450 12345"}
    out.update(over)
    return out


def test_a_misread_name_is_corrected_to_the_documents_spelling():
    out, fixed = verify_identity(ident(name="Lakshmi lyer"), ID_TEXT)  # a capital I read as a lowercase L
    assert out["name"] == "Lakshmi Iyer" and fixed == ["name"]


def test_values_that_already_match_are_untouched_and_not_reported():
    out, fixed = verify_identity(ident(), ID_TEXT)
    assert out == ident() and fixed == []
    out, fixed = verify_identity(ident(name="LAKSHMI IYER", email="Lakshmi.Iyer@Example.com"), ID_TEXT)  # capitals alone are not an error
    assert out["name"] == "LAKSHMI IYER" and out["email"] == "Lakshmi.Iyer@Example.com" and fixed == []


@pytest.mark.parametrize("model_phone", ["+91 9845012345", "9845012345", "+919845012345", "98450-12345"])
def test_a_phone_the_model_tidied_gets_the_documents_own_spelling(model_phone):
    out, fixed = verify_identity(ident(phoneNumber=model_phone), ID_TEXT)
    assert out["phoneNumber"] == "+91 98450 12345" and fixed == ["phoneNumber"]  # keeps the "+91" and the spacing as written


def test_the_documents_spelling_wins_even_without_a_country_code():
    out, fixed = verify_identity(ident(phoneNumber="+91 98450 12345"), "Lakshmi Iyer\nCall 98450 12345 any time\nEXPERIENCE")
    assert out["phoneNumber"] == "98450 12345" and fixed == ["phoneNumber"]


@pytest.mark.parametrize("model_email", ["lakshmi.iyer@examp1e.com", "lakshmi.lyer@example.com", "lakshmi.iyer@exampel.com"])
def test_an_email_a_character_off_is_corrected(model_email):
    out, fixed = verify_identity(ident(email=model_email), ID_TEXT)
    assert out["email"] == "lakshmi.iyer@example.com" and fixed == ["email"]


def test_an_email_split_by_a_line_break_is_still_found():
    text = "Lakshmi Iyer\nlakshmi.iyer@\nexample.com | +91 98450 12345\nEXPERIENCE"
    out, fixed = verify_identity(ident(email="lakshmi.iyer@examp1e.com"), text)
    assert out["email"] == "lakshmi.iyer@example.com" and fixed == ["email"]


def test_a_different_persons_name_is_not_swapped_in():
    out, fixed = verify_identity(ident(name="John Smith"), ID_TEXT)  # "Jane Doe" is in the text too: neither is close enough
    assert out["name"] == "John Smith" and fixed == []


def test_text_with_digits_in_it_is_never_taken_for_a_name():
    # "Priya Verm4" is 0.91 alike "Priya Verma", so it would be swapped in if digits were allowed in a name
    out, fixed = verify_identity(ident(name="Priya Verma"), "Priya Verm4 98765\nEXPERIENCE\nAnalyst at Quillstone Labs")
    assert out["name"] == "Priya Verma" and "name" not in fixed


@pytest.mark.parametrize("field,value", [("email", "someone.else@other.org"), ("phoneNumber", "+91 70000 11111")])
def test_a_value_that_matches_nothing_in_the_text_is_left_alone(field, value):
    out, fixed = verify_identity(ident(**{field: value}), ID_TEXT)
    assert out[field] == value and fixed == []  # nothing is ever invented


def test_with_no_email_or_phone_in_the_text_only_the_name_can_be_checked():
    out, fixed = verify_identity(ident(name="Lakshmi lyer", email="a@b.com"), "Lakshmi Iyer\nno contact details in this document at all")
    assert out["name"] == "Lakshmi Iyer" and out["email"] == "a@b.com" and out["phoneNumber"] == "+91 98450 12345" and fixed == ["name"]


@pytest.mark.parametrize("parsed,text", [
    (None, ID_TEXT), ("not a dict", ID_TEXT), ({"name": "A"}, ""), ({"name": "A"}, None), ({"name": "A"}, "   "),
    ({"name": None, "email": "", "phoneNumber": None}, ID_TEXT), ({"name": 123, "email": ["x"], "phoneNumber": {"a": 1}}, ID_TEXT), ({}, ID_TEXT)])
def test_odd_input_never_crashes_the_identity_check(parsed, text):
    out, fixed = verify_identity(parsed, text)
    assert isinstance(fixed, list)


def test_phone_numbers_finds_every_number_as_written_and_balances_brackets():
    assert list(phone_numbers("a 98765 43210 b +91 91234 56789")) == ["98765 43210", "+91 91234 56789"]
    assert list(phone_numbers("(+91) 98765 43210")) == ["+91 98765 43210"]
    assert list(phone_numbers("(022) 2345 6789")) == ["(022) 2345 6789"]
    assert list(phone_numbers("Jan 2019 - 2023 - 2024 and 2019 2023 98765 43210")) == ["98765 43210"]


def test_a_misread_name_and_a_tidied_phone_are_corrected_through_the_route(client, monkeypatch):
    pdf = _pdf(["Lakshmi Iyer", "lakshmi.iyer@example.com | +91 98450 12345 | Chennai", "EXPERIENCE",
                "Analyst at Quillstone Labs (Jan 2019 - Present)", "EDUCATION", "B.E. in Electronics, Harborview University, 2017", "SKILLS", "Excel, SQL"])
    fake_perfox(monkeypatch, parsed_resume(name="Lakshmi lyer", email="lakshmi.iyer@example.com", phoneNumber="+91 9845012345"))
    r = post(client, "lakshmi.pdf", pdf)
    body = r.json()
    assert r.status_code == 200 and body["name"] == "Lakshmi Iyer" and body["phoneNumber"] == "+91 98450 12345"
    assert body["email"] == "lakshmi.iyer@example.com" and list(body)[-1] == "usage"  # still the flat response, usage last


def test_a_resume_whose_values_are_right_comes_back_unchanged(client, samples, monkeypatch):
    fake_perfox(monkeypatch, parsed_resume())
    body = post(client, "normal_resume.pdf", samples["normal_resume.pdf"].read_bytes()).json()
    assert (body["name"], body["email"], body["phoneNumber"]) == (NAME, EMAIL, PHONE)


def test_the_log_names_the_corrected_fields_but_never_their_values(client, monkeypatch, caplog):
    pdf = _pdf(["Lakshmi Iyer", "lakshmi.iyer@example.com | +91 98450 12345 | Chennai", "EXPERIENCE", "Analyst at Quillstone Labs (Jan 2019 - Present)"])
    fake_perfox(monkeypatch, parsed_resume(name="Lakshmi lyer", email="lakshmi.iyer@example.com", phoneNumber="+91 9845012345"))
    with caplog.at_level(logging.INFO, logger="perfox"):
        assert post(client, "lakshmi.pdf", pdf).status_code == 200
    logged = " ".join(record.getMessage() for record in caplog.records)
    assert "identity corrected: name, phoneNumber" in logged
    for secret in ("Lakshmi", "lyer", "Iyer", "98450", "9845012345", "example.com"):
        assert secret not in logged


# ---------- education entries without an institution (tidy_education) ----------

BE = {"degree": "B.E. in Electronics", "institution": "Harborview University", "graduation_year": 2017}
HSC = {"degree": "HSC", "institution": "", "graduation_year": 2013}


def test_a_school_entry_without_an_institution_is_dropped_when_a_degree_has_one():
    assert tidy_education({"education": [BE, HSC]})["education"] == [BE]
    assert tidy_education({"education": [HSC, BE]})["education"] == [BE]
    assert tidy_education({"education": [BE, dict(HSC, institution="   ")]})["education"] == [BE]  # blanks count as none


def test_a_lone_entry_without_an_institution_is_kept_so_the_app_can_ask():
    assert tidy_education({"education": [HSC]})["education"] == [HSC]
    assert tidy_education({"education": [HSC, dict(HSC, degree="SSC")]})["education"] == [HSC, dict(HSC, degree="SSC")]


def test_entries_that_all_have_an_institution_are_all_kept():
    two = [BE, {"degree": "MBA", "institution": "Lakeview College", "graduation_year": 2020}]
    assert tidy_education({"education": two})["education"] == two


@pytest.mark.parametrize("parsed", [{"education": []}, {}, {"education": None}, {"education": "B.E."}, None, "text"])
def test_education_tidy_up_does_not_crash_on_odd_input(parsed):
    tidy_education(parsed)


def test_a_non_dict_entry_is_dropped_when_a_real_one_has_an_institution():
    assert tidy_education({"education": ["HSC", BE]})["education"] == [BE]


def test_the_route_drops_the_school_entry_and_keeps_the_degree(client, samples, monkeypatch):
    fake_perfox(monkeypatch, parsed_resume(education=[dict(BE), dict(HSC)]))
    body = post(client, "normal_resume.pdf", samples["normal_resume.pdf"].read_bytes()).json()
    assert body["education"] == [BE]


def test_the_route_keeps_a_lone_school_entry(client, samples, monkeypatch):
    fake_perfox(monkeypatch, parsed_resume(education=[dict(HSC)]))
    body = post(client, "normal_resume.pdf", samples["normal_resume.pdf"].read_bytes()).json()
    assert body["education"] == [HSC]


# ---------- the resume prompt ----------

def test_the_resume_prompt_asks_for_every_job_every_skill_and_shown_soft_skills():
    from schema import build_prompt
    prompt = build_prompt(date(2026, 10, 21))
    assert "2026-10-21" in prompt and "{" not in prompt and "}" not in prompt  # the date is filled in and no placeholder is left
    for words in ("5. Work Experience", "EVERY position", "internships", "event-organiser", "6. Skills", "AWS Lambda", "Redis",
                  "7. Soft skills", "Mentorship", "Leadership", "no hobbies"):
        assert words in prompt, words
    for old in ("1. ", "2. ", "3. ", "4. "):  # the four original rules are still there
        assert old in prompt


# ---------- live acceptance: invented resumes already in the repo (RUN_LIVE=1, real Perfox) ----------

FIXTURES = HERE.parents[1] / "tests" / "fixtures"
needs_live = pytest.mark.skipif(os.getenv("RUN_LIVE") != "1", reason="set RUN_LIVE=1 to call the real Perfox endpoint")
needs_keys = pytest.mark.skipif(not (parsers.PERFOX_BASE_URL and parsers.PERFOX_API_KEY), reason="PERFOX_BASE_URL / PERFOX_API_KEY not set")


def live_parse(client, relative: str) -> dict:
    path = FIXTURES / relative
    if not path.exists():
        pytest.skip(f"{relative} is not in this checkout")
    r = client.post("/parsing/resume", files={"pdf_doc": (path.name, path.read_bytes(), PDF)})
    assert r.status_code == 200, r.text[:300]
    return r.json()


@needs_live
@needs_keys
def test_live_a_paragraph_resume_keeps_both_jobs(client):
    jobs = live_parse(client, "junk/files/t03_resume_no_headings_paragraph.pdf")["workExperience"]
    assert len(jobs) >= 2, [j["companyName"] for j in jobs]  # Questlight finds BrightMark Studio and Pixel Hub


@needs_live
@needs_keys
def test_live_cloud_services_and_redis_are_kept(client):
    skills = [s.lower() for s in live_parse(client, "resumes/Rohit_Verma_Python_Backend.pdf")["technical_skills"]]
    assert any("redis" in s for s in skills) and any("aws" in s for s in skills), skills


@needs_live
@needs_keys
@pytest.mark.parametrize("relative", ["resumes/Ananya_Krishnan_Linux_DevOps.pdf", "junk/files/r01_standard.pdf"])
def test_live_soft_skills_are_listed(client, relative):
    assert len(live_parse(client, relative)["soft_skills"]) >= 1


@needs_live
@needs_keys
def test_live_the_name_is_not_misread(client):
    name = live_parse(client, "junk/files/r11_odd_heads.pdf")["name"].lower()
    assert "iyer" in name and "lyer" not in name, name


@needs_live
@needs_keys
def test_live_no_education_entry_is_left_without_an_institution(client):
    education = live_parse(client, "junk/files/r07_fresher.pdf")["education"]
    assert education and all(str(e["institution"]).strip() for e in education), education


@needs_live
@needs_keys
def test_live_the_phone_is_kept_as_the_document_writes_it(client):
    # The document says 9876501234 with no country code. Questlight's parser answers "+91 9876501234"; this service keeps what is written.
    assert live_parse(client, "junk/files/r08_nontech_sales.pdf")["phoneNumber"] == "9876501234"


# ---------- masking: what counts as contact detail (masking.find_sensitive) ----------

MARKED = ["a.b@example.com", "x@sub.example.co.in", "c@x.io", "Meera.Nair@Example.COM", "meera.nair  @example.com", "+91-9876543210", "98765 43210",
          "020 2345 6789", "(022) 2345 6789", "12345678901", "9876543210123", "98450 12345 ext 22", "linkedin.com/in/x",
          "https://www.linkedin.com/in/meera-nair-fake", "in.linkedin.com/in/x", "github.com/x", "https://github.com/x/some-repo", "www.github.com/x"]
NOT_MARKED = ["https://www.linkedin.com/company/x", "twitter.com/x", "instagram.com/x", "gitlab.com/x", "stackoverflow.com/users/1/x", "https://my-portfolio.example.org",
              "14 March 1995", "14/03/1995", "Jan 2019 - Dec 2023", "2019 - 2023 - 2024", "411001", "98765 4321", "Meera Nair, Pune", "call me @ 5.30pm",
              "mygithub.com/x", "github.com", "Employee Ref 12345 and 67890"]


@pytest.mark.parametrize("value", MARKED)
def test_contact_details_are_found_whole(value):
    text = f"Contact: {value} end"
    assert [text[a:b] for a, b in masking.find_sensitive(text)] == [value]


@pytest.mark.parametrize("value", NOT_MARKED)
def test_everything_else_is_left_alone(value):
    assert masking.find_sensitive(f"Note {value} here") == []


def test_two_on_a_line_trailing_punctuation_and_a_sentence_email():
    text = "a@b.com and c@d.org, mail a@b.com. Contact HR at hr@quillstone.example for a reference. See github.com/x/y, thanks"
    assert [text[a:b] for a, b in masking.find_sensitive(text)] == ["a@b.com", "c@d.org", "a@b.com", "hr@quillstone.example", "github.com/x/y"]


def test_phone_spans_say_where_each_number_sits():
    text = "Call 98765 43210 or (+91) 91234 56789 now"
    from fix_normlization import phone_spans
    assert [text[a:b] for a, b in phone_spans(text)] == ["98765 43210", "+91) 91234 56789"]


# ---------- masking a PDF in place (masking.mask_pdf) ----------

MASK_LINES = [
    "Meera Nair", "Pune, Maharashtra 411001",
    "Email: meera.nair@example.com   Alt: meera.alt@example.co.in", "Phone: +91 91234 56789   Home: 020 2345 6789 ext 22",
    "LinkedIn: https://www.linkedin.com/in/meera-nair-fake  GitHub: github.com/meera-fake/my-repo",
    "Company: https://www.linkedin.com/company/northwind-supplies  Blog: https://meera.example.org/blog  X: twitter.com/meera_nair",
    "Date of Birth: 14 March 1995    Postal: 411001    Years: 2019 - 2023",
    "Data Analyst, Quillstone Labs, Pune (Jan 2021 - Present)", "Skills: Excel, SQL, Power BI",
]
MASK_REMOVED = ["meera.nair@example.com", "meera.alt@example.co.in", "91234 56789", "020 2345 6789", "ext 22", "linkedin.com/in/meera-nair-fake", "github.com/meera-fake"]
MASK_KEPT = ["Meera Nair", "Pune, Maharashtra 411001", "linkedin.com/company/northwind-supplies", "meera.example.org/blog", "twitter.com/meera_nair",
             "14 March 1995", "2019 - 2023", "Postal: 411001", "Data Analyst, Quillstone Labs, Pune (Jan 2021 - Present)", "Skills: Excel, SQL, Power BI"]
MASK_COUNT = 6  # two e-mails, two phone numbers (the extension goes with its number), one LinkedIn profile, one GitHub link


def _mask_pdf_bytes(font: str = "Helvetica", lines: list[str] | None = None, **info) -> bytes:
    import reportlab
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    if info.get("author"):
        c.setAuthor(info["author"])
    if info.get("title"):
        c.setTitle(info["title"])
    if font == "Vera":  # reportlab's bundled TrueType font: PDF writers embed only the letters used (a "subset" font), as Word and Google Docs do
        pdfmetrics.registerFont(TTFont("VeraT", os.path.join(os.path.dirname(reportlab.__file__), "fonts", "Vera.ttf")))
        font = "VeraT"
    c.setFont(font, 10)
    y = 800
    for line in lines or MASK_LINES:
        c.drawString(40, y, line)
        y -= 22
    c.save()
    return buf.getvalue()


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    return "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(data)).pages)


def _pdf_raw(data: bytes) -> bytes:
    """The file's bytes plus the inflated content streams, so a value hidden inside a compressed stream would still be found."""
    import re
    import zlib
    blob = data
    for match in re.finditer(rb"stream\r?\n(.*?)\r?\nendstream", data, re.S):
        try:
            blob += zlib.decompress(match.group(1))
        except zlib.error:
            pass
    return blob


def _pages(data: bytes) -> int:
    from pypdf import PdfReader
    return len(PdfReader(io.BytesIO(data)).pages)


@pytest.mark.parametrize("font", ["Helvetica", "Vera"])
def test_contact_details_are_removed_in_place_and_everything_else_stays(font):
    source = _mask_pdf_bytes(font)
    masked, count, warnings = masking.mask_pdf(source)
    text, raw_bytes = _pdf_text(masked), _pdf_raw(masked)
    assert count == MASK_COUNT and warnings == [] and masked[:5] == b"%PDF-"
    for value in MASK_REMOVED:
        assert value not in text and value.encode() not in raw_bytes, f"still there: {value}"
    for value in MASK_KEPT:
        assert value in text, f"lost: {value}"
    assert "Email:" in text and "Phone:" in text  # the labels stay, with a gap after them, as in Questlight's
    assert _pages(masked) == _pages(source)


def _split_pieces_pdf() -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setFont("Helvetica", 11)

    def row(y: int, words: list[tuple[str, int]]):
        x = 40
        for word, gap in words:
            c.drawString(x, y, word)
            x += c.stringWidth(word, "Helvetica", 11) + gap
    row(790, [("Email:", 6), ("meera.nair", 0), ("@example.com", 0)])               # the pieces touch: no space between them at all
    row(766, [("Mail:", 6), ("priya.verma ", 0), ("@example.org", 0)])              # the first piece ends with a space
    row(742, [("Phone:", 6), ("+91", 5), ("91234", 5), ("56789", 0)])               # one number in three pieces
    row(718, [("Analyst,", 5), ("Quillstone", 5), ("Labs", 0)])
    c.save()
    return buf.getvalue()


def test_values_drawn_as_separate_words_are_masked_too():
    masked, count, warnings = masking.mask_pdf(_split_pieces_pdf())
    text, raw_bytes = _pdf_text(masked), _pdf_raw(masked)
    assert count == 3 and warnings == []
    for bit in (b"meera.nair", b"priya.verma", b"91234", b"56789", b"example"):
        assert bit not in raw_bytes, bit
    assert "Analyst, Quillstone Labs" in text and "Email:" in text and "Mail:" in text and "Phone:" in text


def _styled_pdf() -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setFillColor(colors.HexColor("#1F3A5F"))
    c.rect(0, 760, 595, 82, fill=1, stroke=0)
    c.setFillColor(colors.white)
    c.setFont("Times-BoldItalic", 26)
    c.drawString(40, 792, "Meera Nair")
    c.setFillColor(colors.black)
    c.setFont("Courier", 10)
    c.drawString(40, 735, "Email: meera.nair@example.com   Phone: +91 91234 56789")  # the only line with contact details
    c.setStrokeColor(colors.red)
    c.setLineWidth(3)
    c.line(40, 718, 555, 718)
    c.setFont("Helvetica-Oblique", 14)
    c.drawString(300, 650, "Slanted text at an odd place")
    c.saveState()
    c.translate(120, 520)
    c.rotate(25)
    c.setFont("Helvetica", 12)
    c.drawString(0, 0, "Rotated line of text")
    c.restoreState()
    c.setFillColor(colors.HexColor("#2E7D32"))
    c.circle(480, 600, 28, fill=1, stroke=0)
    c.setFillColor(colors.black)
    c.setFont("Helvetica", 11)
    c.drawString(40, 480, "Data Analyst, Quillstone Labs, Pune (Jan 2021 - Present)")
    c.save()
    return buf.getvalue()


def test_on_a_styled_page_only_the_contact_line_changes():
    import pypdfium2 as pdfium
    from PIL import ImageChops
    source = _styled_pdf()
    masked, count, warnings = masking.mask_pdf(source)
    assert count == 2 and warnings == []

    def render(data: bytes):
        return pdfium.PdfDocument(data)[0].render(scale=2).to_pil().convert("RGB")
    diff = ImageChops.difference(render(source), render(masked)).convert("L").point(lambda p: 255 if p > 40 else 0)
    top, bottom = (842 - 750) * 2, (842 - 727) * 2  # the band of the contact line, in pixels
    assert diff.crop((0, top, diff.width, bottom)).getbbox() is not None  # the contact line did change
    diff.paste(0, (0, top, diff.width, bottom))
    assert diff.getbbox() is None  # and nothing else did: the banner, title, red rule, slanted and rotated text, circle and the job line are identical
    text = _pdf_text(masked)
    assert "Meera Nair" in text and "Rotated line of text" in text and "meera.nair" not in text


def test_every_link_is_removed_but_form_fields_stay():
    from pypdf import PdfReader
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setFont("Helvetica", 11)
    y = 800
    for label, url in [("mail me", "mailto:meera.nair@example.com"), ("call me", "tel:+919123456789"), ("my linkedin", "https://www.linkedin.com/in/meera-nair-fake"),
                       ("my twitter", "https://twitter.com/meera_nair"), ("company", "https://www.linkedin.com/company/x")]:
        c.drawString(50, y, label)
        c.linkURL(url, (50, y - 3, 130, y + 12))
        y -= 24
    c.acroForm.textfield(name="note", tooltip="note", x=50, y=600, width=100, height=20, value="v")
    c.save()

    def kinds(data: bytes) -> list[str]:
        return [str(a.get_object().get("/Subtype")) for a in (PdfReader(io.BytesIO(data)).pages[0].get("/Annots") or [])]
    source = buf.getvalue()
    assert kinds(source).count("/Link") == 5
    masked, count, warnings = masking.mask_pdf(source)
    assert kinds(masked) == ["/Widget"] and count == 0 and "mail me" in _pdf_text(masked)


def test_the_pdfs_metadata_is_not_touched_like_questlights():
    from pypdf import PdfReader
    source = _mask_pdf_bytes(author="Meera Nair meera.nair@example.com", title="Resume of Meera Nair")
    meta = PdfReader(io.BytesIO(masking.mask_pdf(source)[0])).metadata
    # The author and title are left as they are, as Questlight's mask does; so do we, on purpose.
    assert meta.author == "Meera Nair meera.nair@example.com" and meta.title == "Resume of Meera Nair"


def test_a_scan_comes_back_with_a_warning_and_nothing_masked():
    masked, count, warnings = masking.mask_pdf(_scanned_pdf())
    assert count == 0 and warnings == ["no text layer"] and masked[:5] == b"%PDF-" and _pages(masked) == 1


def test_a_pdf_with_nothing_to_mask_keeps_its_text_and_reports_zero():
    lines = ["Meera Nair", "Analyst at Quillstone Labs", "Skills: Excel, SQL and Power BI for reporting.", "Postal: 411001  Years: 2019 - 2023"]
    masked, count, warnings = masking.mask_pdf(_mask_pdf_bytes(lines=lines))
    assert count == 0 and warnings == [] and _pdf_text(masked).split() == " ".join(lines).split()


def test_a_value_that_could_not_be_removed_is_reported_not_hidden(monkeypatch):
    monkeypatch.setattr(masking, "_blank_line", lambda line: 0)  # simulate masking that fails to edit anything
    masked, count, warnings = masking.mask_pdf(_mask_pdf_bytes())
    assert count == 0 and warnings == ["leftovers"] and masked[:5] == b"%PDF-"  # still returned, as Questlight would, but flagged


def _encrypted_pdf() -> bytes:
    from pypdf import PdfReader, PdfWriter
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(_mask_pdf_bytes())))
    writer.encrypt("secret")
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


@pytest.mark.parametrize("data", [b"", b"this is not a pdf at all " * 20, b"%PDF-1.4 but then nothing that makes sense", _encrypted_pdf()])
def test_bytes_that_are_not_an_openable_pdf_are_refused(data):
    with pytest.raises(masking.CannotOpen):
        masking.mask_pdf(data)


# ---------- masking: the route POST /parsing/mask ----------

def post_mask(client, data: bytes, name: str = "meera.pdf"):
    return client.post("/parsing/mask", files={"file": (name, data, PDF)})


def test_the_mask_route_returns_the_masked_pdf_as_a_download(client):
    r = post_mask(client, _mask_pdf_bytes())
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf" and r.content[:5] == b"%PDF-"
    assert r.headers["content-disposition"] == "attachment; filename=meera_redacted.pdf"
    assert r.headers["x-masked-items"] == str(MASK_COUNT) and r.headers["x-request-id"] and "x-masking-warning" not in r.headers
    assert "meera.nair@example.com" not in _pdf_text(r.content) and "Meera Nair" in _pdf_text(r.content)


def test_the_download_name_handles_spaces_and_non_ascii_names(client):
    spaced = post_mask(client, _mask_pdf_bytes(), "Meera Nair resume.pdf").headers["content-disposition"]
    assert spaced == "attachment; filename=\"Meera Nair resume_redacted.pdf\"; filename*=UTF-8''Meera%20Nair%20resume_redacted.pdf"
    accented = post_mask(client, _mask_pdf_bytes(), "Mérā.pdf").headers["content-disposition"]
    assert accented == "attachment; filename=\"Mr_redacted.pdf\"; filename*=UTF-8''M%C3%A9r%C4%81_redacted.pdf"


def test_a_scan_is_still_a_200_with_a_warning_header(client):
    r = post_mask(client, _scanned_pdf())
    assert r.status_code == 200 and r.content[:5] == b"%PDF-" and r.headers["x-masked-items"] == "0" and r.headers["x-masking-warning"] == "no text layer"


def test_no_file_is_a_400_in_questlights_words(client):
    r = client.post("/parsing/mask")
    err = r.json()["error"]
    assert r.status_code == 400 and err["code"] == 400 and err["error_type"] == "INVALID_REQUEST" and r.json()["success"] is False
    assert err["message"] == "No file was uploaded. Please choose a file and try again."


@pytest.mark.parametrize("name,data", [("meera.pdf", b""), ("meera.pdf", b"not a pdf " * 10), ("meera.docx", _mask_pdf_bytes()), ("meera.txt", b"%PDF-1.4 text"),
                                       ("meera.pdf", _encrypted_pdf())])
def test_anything_that_is_not_a_readable_pdf_is_a_400_unsupported(client, name, data):
    r = post_mask(client, data, name)
    err = r.json()["error"]
    assert r.status_code == 400 and err["error_type"] == "UNSUPPORTED_FILE_TYPE" and err["code"] == 400
    assert err["message"] == "This file type isn't supported for redaction. Please upload a PDF."


def test_a_file_over_the_limit_is_a_400(client):
    r = post_mask(client, b"%PDF-1.4\n" + b"0" * parsers.MAX_FILE_BYTES)
    assert r.status_code == 400 and r.json()["error"]["error_type"] == "INVALID_REQUEST" and "too large" in r.json()["error"]["message"]


def test_an_unexpected_failure_is_a_500_with_no_file_and_no_content_in_the_log(client, monkeypatch, caplog):
    def explode(data):
        raise RuntimeError("boom with meera.nair@example.com inside")
    monkeypatch.setattr(masking, "mask_pdf", explode)
    with caplog.at_level(logging.INFO, logger="perfox"):
        r = post_mask(client, _mask_pdf_bytes())
    assert r.status_code == 500 and r.json()["error"]["error_type"] == "MASKING_FAILED" and r.content[:5] != b"%PDF-"
    logged = " ".join(record.getMessage() for record in caplog.records)
    assert "RuntimeError" in logged and "meera.nair" not in logged  # the kind of error is logged, never what it said


def test_masking_never_logs_what_it_removed(client, caplog):
    with caplog.at_level(logging.INFO, logger="perfox"):
        assert post_mask(client, _mask_pdf_bytes()).status_code == 200
    logged = " ".join(record.getMessage() for record in caplog.records)
    assert f"items={MASK_COUNT}" in logged
    for secret in ("meera.nair", "meera.alt", "91234", "020 2345", "linkedin", "github", "Meera", "Quillstone"):
        assert secret not in logged, secret


def test_masking_makes_no_perfox_call_and_costs_no_tokens(client, monkeypatch):
    async def must_not_be_called(*args, **kwargs):
        raise AssertionError("masking must not call Perfox")
    monkeypatch.setattr(parsers, "call_perfox", must_not_be_called)
    r = post_mask(client, _mask_pdf_bytes())
    assert r.status_code == 200 and "usage" not in r.headers


def test_contract_the_apps_own_mask_call_accepts_our_answer(monkeypatch):
    """app/intake/parser_questlight.mask() posts the file to {PARSING_BASE_URL}/mask and accepts the answer only if it is a 200 PDF."""
    import asyncio
    import base64
    import httpx
    from app import settings
    from app.intake import parser_questlight
    monkeypatch.setattr(settings, "PARSING_BASE_URL", "http://svc/parsing")

    async def run():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=parsers.app), base_url="http://svc", timeout=60) as service:
            return await parser_questlight.mask(service, "resume.pdf", _mask_pdf_bytes(), ".pdf")
    encoded, warning = asyncio.run(run())
    assert warning is None
    masked = base64.b64decode(encoded)
    assert masked[:5] == b"%PDF-" and "meera.nair@example.com" not in _pdf_text(masked)


# ---------- the token check (SERVICE_TOKEN) ----------

TOKEN = "unit-test-token-not-a-real-secret"
PARSING_ROUTES = ["/parsing/resume", "/parsing/jd", "/parsing/mask"]


def bearer(value: str = TOKEN, scheme: str = "Bearer") -> dict:
    return {"Authorization": f"{scheme} {value}"}


def test_without_the_setting_the_service_is_open(client, monkeypatch):
    monkeypatch.delenv("SERVICE_TOKEN", raising=False)
    assert client.post("/parsing/mask", files={"file": ("meera.pdf", _mask_pdf_bytes(), PDF)}).status_code == 200


def test_an_empty_setting_also_means_open(client, monkeypatch):
    monkeypatch.setenv("SERVICE_TOKEN", "")
    assert client.post("/parsing/mask", files={"file": ("meera.pdf", _mask_pdf_bytes(), PDF)}).status_code == 200


@pytest.mark.parametrize("path", PARSING_ROUTES)
def test_every_parsing_route_needs_the_token_when_it_is_set(client, monkeypatch, path):
    monkeypatch.setenv("SERVICE_TOKEN", TOKEN)
    r = client.post(path)  # no header, and no file: the check comes before everything else
    err = r.json()["error"]
    assert r.status_code == 401 and err["code"] == 401 and err["error_type"] == "UNAUTHORIZED" and err["request_id"] and r.json()["success"] is False
    assert r.headers["www-authenticate"] == "Bearer" and r.headers["x-request-id"] == err["request_id"] and TOKEN not in r.text


@pytest.mark.parametrize("headers", [bearer("wrong"), bearer(""), bearer(TOKEN[:-1]), bearer(TOKEN + "x"), bearer(TOKEN, "Basic"), bearer(TOKEN.upper()),
                                     {"Authorization": TOKEN}, {"Authorization": "Bearer"}, {"X-Api-Key": TOKEN}])
def test_a_wrong_or_malformed_token_is_refused(client, monkeypatch, headers):
    monkeypatch.setenv("SERVICE_TOKEN", TOKEN)
    r = client.post("/parsing/mask", files={"file": ("m.pdf", _mask_pdf_bytes(), PDF)}, headers=headers)
    assert r.status_code == 401 and r.content[:5] != b"%PDF-"


@pytest.mark.parametrize("scheme", ["Bearer", "bearer", "BEARER"])
def test_the_right_token_gets_through_on_every_route(client, monkeypatch, scheme):
    monkeypatch.setenv("SERVICE_TOKEN", TOKEN)
    headers = bearer(scheme=scheme)
    assert client.post("/parsing/mask", files={"file": ("m.pdf", _mask_pdf_bytes(), PDF)}, headers=headers).status_code == 200
    fake_perfox(monkeypatch, parsed_jd())
    jd_text = "Senior QA Automation Engineer wanted: four to eight years with Java and Selenium, based in Bengaluru, 12-18 LPA."
    assert client.post("/parsing/jd", data={"jd_text": jd_text}, headers=headers).status_code == 200
    fake_perfox(monkeypatch, parsed_resume())
    assert client.post("/parsing/resume", files={"pdf_doc": ("normal.pdf", _pdf(NORMAL_LINES), PDF)}, headers=headers).status_code == 200


def test_health_and_the_docs_stay_open(client, monkeypatch):
    monkeypatch.setenv("SERVICE_TOKEN", TOKEN)
    assert client.get("/health").status_code == 200 and client.get("/docs").status_code == 200 and client.get("/openapi.json").status_code == 200


def test_a_refused_call_costs_nothing_and_logs_no_token(client, monkeypatch, caplog):
    monkeypatch.setenv("SERVICE_TOKEN", TOKEN)
    with caplog.at_level(logging.INFO, logger="perfox"):
        r = client.post("/parsing/resume", files={"pdf_doc": ("normal.pdf", _pdf(NORMAL_LINES), PDF)}, headers=bearer("a-wrong-token-value"))
    assert r.status_code == 401  # the autouse fixture makes any call to Perfox fail loudly, so none was made
    logged = " ".join(record.getMessage() for record in caplog.records)
    assert "UNAUTHORIZED" in logged and TOKEN not in logged and "a-wrong-token-value" not in logged


# ---------- packaging for Vercel ----------

SERVICE_DIR = HERE.parent
NOT_UPLOADED = {"updates.md", "BUILD_SPEC.md"}


def test_pyproject_and_requirements_list_the_same_dependencies():
    import tomllib
    dependencies = tomllib.loads((SERVICE_DIR / "pyproject.toml").read_text(encoding="utf-8"))["project"]["dependencies"]
    requirements = [ln.strip() for ln in (SERVICE_DIR / "requirements.txt").read_text(encoding="utf-8").splitlines() if ln.strip() and not ln.startswith("#")]
    assert sorted(dependencies) == sorted(requirements)
    assert {"pypdf>=6", "pypdfium2>=5", "httpx>=0.28,<1"} <= set(dependencies)


def test_the_vercel_settings_point_at_files_that_exist():
    import json
    import tomllib
    module, attribute = tomllib.loads((SERVICE_DIR / "pyproject.toml").read_text(encoding="utf-8"))["tool"]["vercel"]["entrypoint"].split(":")
    assert (SERVICE_DIR / f"{module}.py").exists() and hasattr(parsers, attribute)
    functions = json.loads((SERVICE_DIR / "vercel.json").read_text(encoding="utf-8"))["functions"]
    assert all((SERVICE_DIR / name).exists() for name in functions) and functions["parsers.py"]["maxDuration"] >= parsers.TIMEOUT_S
    ignored = (SERVICE_DIR / ".vercelignore").read_text(encoding="utf-8").split()
    assert {".env", "tests", "updates.md", "BUILD_SPEC.md"} <= set(ignored)
    assert (SERVICE_DIR / ".python-version").read_text(encoding="utf-8").strip() == "3.12"


def test_the_service_runs_on_its_own_without_the_main_app(tmp_path):
    """Copies the files Vercel would upload into an empty folder and runs the service there, with nothing of the main app reachable."""
    import shutil
    import subprocess
    for path in SERVICE_DIR.iterdir():
        if path.is_file() and path.suffix in {".py", ".md", ".toml", ".json", ".txt"} and path.name not in NOT_UPLOADED:
            shutil.copy(path, tmp_path / path.name)
    (tmp_path / "sample.pdf").write_bytes(_pdf(NORMAL_LINES))
    script = (
        "import sys\n"
        "import parsers\n"
        "assert 'app' not in sys.modules, 'the service pulled in the main app'\n"
        "from fastapi.testclient import TestClient\n"
        "async def fake(*args):\n"
        f"    return {{'name': {NAME!r}, 'email': {EMAIL!r}, 'phoneNumber': {PHONE!r}}}, {{'input_tokens': 1, 'output_tokens': 1}}, None\n"
        "parsers.call_perfox = fake\n"
        "client, data = TestClient(parsers.app), open('sample.pdf', 'rb').read()\n"
        "assert client.get('/health').status_code == 200\n"
        "r = client.post('/parsing/resume', files={'pdf_doc': ('r.pdf', data, 'application/pdf')}); assert r.status_code == 200 and r.json()['email'] == " + repr(EMAIL) + ", r.text[:200]\n"
        "r = client.post('/parsing/jd', data={'jd_text': 'A Java developer with five years of experience in Spring Boot and microservices.'}); assert r.status_code == 200, r.text[:200]\n"
        "r = client.post('/parsing/mask', files={'file': ('r.pdf', data, 'application/pdf')}); assert r.status_code == 200 and r.content[:5] == b'%PDF-' and r.headers['x-masked-items'] == '2', r.headers\n"
        "assert 'app' not in sys.modules\n"
        "print('standalone ok')\n"
    )
    env = {key: value for key, value in os.environ.items() if key not in {"PYTHONPATH", "SERVICE_TOKEN"}}
    env.update(PYTHONDONTWRITEBYTECODE="1", PERFOX_BASE_URL="", PERFOX_API_KEY="")
    run = subprocess.run([sys.executable, "-c", script], cwd=tmp_path, env=env, capture_output=True, text=True, timeout=120)
    assert run.returncode == 0 and "standalone ok" in run.stdout, (run.stdout + run.stderr)[-800:]
